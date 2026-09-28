"""Tests for parser.py (FR-2, E1).

Covers:
  - parse_pdf: page count, 1-indexed page numbers, non-empty text.
  - parse_pptx: slide count, 1-indexed slide numbers, non-empty text.
  - Both parsers skip blank pages/slides.
"""

from __future__ import annotations

from pathlib import Path

from coursemind_ingest.parser import parse_pdf, parse_pptx

# ---------------------------------------------------------------------------
# PDF tests
# ---------------------------------------------------------------------------


class TestParsePdf:
    def test_page_count(self, sample_pdf: Path) -> None:
        """sample.pdf has 2 text pages; parser should return both."""
        pages = parse_pdf(sample_pdf)
        assert len(pages) == 2, f"Expected 2 pages, got {len(pages)}"

    def test_page_numbers_are_positive_and_one_indexed(self, sample_pdf: Path) -> None:
        pages = parse_pdf(sample_pdf)
        numbers = [p for p, _ in pages]
        assert all(n >= 1 for n in numbers), f"Non-positive page number: {numbers}"
        assert numbers[0] == 1, "First page should be page 1"

    def test_text_is_non_empty(self, sample_pdf: Path) -> None:
        pages = parse_pdf(sample_pdf)
        for page_num, text in pages:
            assert text.strip(), f"Empty text on page {page_num}"

    def test_text_contains_expected_content(self, sample_pdf: Path) -> None:
        """Sanity-check that text extraction is working, not just returning whitespace."""
        pages = parse_pdf(sample_pdf)
        all_text = " ".join(t for _, t in pages).lower()
        assert "machine learning" in all_text

    def test_returns_list_of_tuples(self, sample_pdf: Path) -> None:
        pages = parse_pdf(sample_pdf)
        assert isinstance(pages, list)
        for item in pages:
            assert isinstance(item, tuple)
            assert len(item) == 2
            assert isinstance(item[0], int)
            assert isinstance(item[1], str)

    def test_blank_page_skipped(self, tmp_path: Path) -> None:
        """A PDF with a blank page should return only the non-blank pages."""
        import pymupdf as fitz

        pdf_path = tmp_path / "blank.pdf"
        doc = fitz.open()
        doc.new_page()  # blank page
        page2 = doc.new_page()
        page2.insert_text((50, 50), "Hello world")
        doc.save(str(pdf_path))
        doc.close()

        pages = parse_pdf(pdf_path)
        assert len(pages) == 1, f"Expected 1 non-blank page, got {len(pages)}"
        assert pages[0][0] == 2, "Non-blank page should be page 2"


# ---------------------------------------------------------------------------
# PPTX tests
# ---------------------------------------------------------------------------


class TestParsePptx:
    def test_slide_count(self, sample_pptx: Path) -> None:
        """sample.pptx has 3 slides; parser should return all 3."""
        slides = parse_pptx(sample_pptx)
        assert len(slides) == 3, f"Expected 3 slides, got {len(slides)}"

    def test_slide_numbers_are_positive_and_one_indexed(self, sample_pptx: Path) -> None:
        slides = parse_pptx(sample_pptx)
        numbers = [s for s, _ in slides]
        assert all(n >= 1 for n in numbers), f"Non-positive slide number: {numbers}"
        assert numbers[0] == 1, "First slide should be slide 1"

    def test_text_is_non_empty(self, sample_pptx: Path) -> None:
        slides = parse_pptx(sample_pptx)
        for slide_num, text in slides:
            assert text.strip(), f"Empty text on slide {slide_num}"

    def test_text_contains_expected_content(self, sample_pptx: Path) -> None:
        slides = parse_pptx(sample_pptx)
        all_text = " ".join(t for _, t in slides).lower()
        assert "machine learning" in all_text

    def test_returns_list_of_tuples(self, sample_pptx: Path) -> None:
        slides = parse_pptx(sample_pptx)
        assert isinstance(slides, list)
        for item in slides:
            assert isinstance(item, tuple)
            assert len(item) == 2
            assert isinstance(item[0], int)
            assert isinstance(item[1], str)

    def test_blank_slide_skipped(self, tmp_path: Path) -> None:
        """A PPTX with a blank slide should return only the non-blank slides."""
        from pptx import Presentation

        pptx_path = tmp_path / "blank.pptx"
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]  # blank layout
        prs.slides.add_slide(blank_layout)  # truly blank slide

        content_layout = prs.slide_layouts[1]
        slide2 = prs.slides.add_slide(content_layout)
        slide2.shapes.title.text = "Not blank"
        slide2.placeholders[1].text = "Some content here"
        prs.save(str(pptx_path))

        slides = parse_pptx(pptx_path)
        assert len(slides) == 1, f"Expected 1 non-blank slide, got {len(slides)}"
        assert slides[0][0] == 2, "Non-blank slide should be slide 2"
